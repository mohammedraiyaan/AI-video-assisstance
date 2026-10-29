"""Build an editable AAT PowerPoint from project documentation and real results."""

import argparse
import csv
from pathlib import Path


def build_presentation(
    output_path: Path = Path("docs/aat_presentation.pptx"),
    results_path: Path = Path("results/results.csv"),
) -> Path:
    """Create a concise slide deck; absent evaluation data is labeled as pending."""
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError(
            "Install python-pptx to build the AAT presentation"
        ) from exc

    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)

    def add_title(slide, title: str, subtitle: str | None = None) -> None:
        title_box = slide.shapes.title
        title_box.text = title
        title_box.text_frame.paragraphs[0].font.size = Pt(30)
        title_box.text_frame.paragraphs[0].font.bold = True
        title_box.text_frame.paragraphs[0].font.color.rgb = RGBColor(54, 35, 93)
        if subtitle:
            box = slide.shapes.add_textbox(Inches(0.7), Inches(1.25), Inches(12), Inches(0.5))
            paragraph = box.text_frame.paragraphs[0]
            paragraph.text = subtitle
            paragraph.font.size = Pt(16)

    def add_bullets(slide, bullets: list[str], top: float = 1.6) -> None:
        box = slide.shapes.add_textbox(
            Inches(0.8), Inches(top), Inches(11.9), Inches(5.3)
        )
        frame = box.text_frame
        frame.word_wrap = True
        for index, text in enumerate(bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = text
            paragraph.level = 0
            paragraph.font.size = Pt(20)
            paragraph.space_after = Pt(13)

    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "RAG for an AI Video Summarizer"
    slide.shapes.title.text_frame.paragraphs[0].font.color.rgb = RGBColor(54, 35, 93)
    slide.placeholders[1].text = (
        "AAT project | Timestamped transcription, retrieval, grounded generation, and evaluation"
    )

    slide = deck.slides.add_slide(deck.slide_layouts[1])
    add_title(slide, "Architecture components")
    add_bullets(
        slide,
        [
            "Input and audio extraction: YouTube URLs or local audio/video",
            "Transcription: local Whisper for English; Sarvam for Hindi/Hinglish",
            "Timestamp-aware chunking and transcript/vector caching",
            "BM25 sparse index and HuggingFace + Chroma dense index",
            "Hybrid RRF plus cross-encoder reranking",
            "Gemini generation and judging; Streamlit UI with citations and Markdown export",
            "Evaluation runner, CSV/PNG results, error analysis, and report outputs",
        ],
    )

    slide = deck.slides.add_slide(deck.slide_layouts[1])
    add_title(slide, "Data flow")
    add_bullets(
        slide,
        [
            "1. Accept a YouTube URL or local media file.",
            "2. Validate the source and extract/resample audio.",
            "3. Detect language and transcribe into timestamped segments.",
            "4. Cache the transcript by source and language.",
            "5. Create configurable, overlapping timestamp-aware chunks.",
            "6. Embed/index chunks for sparse and/or dense retrieval.",
            "7. Retrieve with baseline, BM25, dense, or hybrid RRF + reranking.",
            "8. Generate a summary or grounded answer with clickable citations.",
            "9. Record evaluation metrics and present results in the app/report.",
        ],
        top=1.45,
    )

    slide = deck.slides.add_slide(deck.slide_layouts[5])
    add_title(slide, "Approaches compared")
    table = slide.shapes.add_table(5, 3, Inches(0.5), Inches(1.5), Inches(12.3), Inches(4.5)).table
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(4.8)
    table.columns[2].width = Inches(5.5)
    comparison = [
        ("Approach", "Method", "Trade-off"),
        ("Baseline", "Bounded full transcript", "Simple reference; large context and truncation risk"),
        ("Sparse", "BM25 lexical ranking", "Fast; weak on paraphrase"),
        ("Dense", "HuggingFace embeddings + Chroma", "Semantic matching; embedding/index cost"),
        ("Hybrid", "BM25 + dense RRF + cross-encoder", "More compute/latency; combines lexical and semantic signals"),
    ]
    for row_index, values in enumerate(comparison):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.text = value
            cell.text_frame.paragraphs[0].font.size = Pt(16)

    slide = deck.slides.add_slide(deck.slide_layouts[1])
    add_title(slide, "Evaluation metrics")
    add_bullets(
        slide,
        [
            "Retrieval: Recall@3, Recall@5, MRR, hit rate",
            "Answer quality: Gemini-judged faithfulness/relevance, ROUGE-L, BERTScore",
            "Summary quality: reference ROUGE-L/BERTScore when supplied; compression ratio",
            "System: end-to-end query latency, transcription time, estimated tokens and configured-rate cost",
            "Ablation: dense retrieval with chunk sizes 256/512/1024 and overlap 0/64",
            "No scores are shown until the evaluation is run with real videos and reviewed references.",
        ],
    )

    slide = deck.slides.add_slide(deck.slide_layouts[1])
    add_title(slide, "Four planned test cases")
    add_bullets(
        slide,
        [
            "English lecture — replace placeholder in videos.yaml",
            "Hindi or Hinglish talk — replace placeholder in videos.yaml",
            "Short video under five minutes — replace placeholder in videos.yaml",
            "Long video over 30 minutes or noisy audio — replace placeholder in videos.yaml",
            "Hand-edit generated evaluation/data/<video_id>_qa.json references before reporting scores.",
        ],
    )

    slide = deck.slides.add_slide(deck.slide_layouts[5])
    add_title(slide, "Measured results")
    if results_path.is_file():
        with results_path.open(encoding="utf-8", newline="") as file:
            rows = list(csv.DictReader(file))
        headers = [
            "approach",
            "recall@3",
            "recall@5",
            "mrr",
            "hit_rate",
            "faithfulness",
            "answer_relevance",
            "end_to_end_latency_seconds",
        ]
        table = slide.shapes.add_table(
            len(rows) + 1,
            len(headers),
            Inches(0.25),
            Inches(1.65),
            Inches(12.85),
            Inches(min(4.8, 0.55 * (len(rows) + 1))),
        ).table
        for column_index, header in enumerate(headers):
            table.cell(0, column_index).text = header
        for row_index, row in enumerate(rows, 1):
            for column_index, header in enumerate(headers):
                value = row.get(header, "")
                try:
                    value = f"{float(value):.3f}" if value else "--"
                except ValueError:
                    value = str(value)
                table.cell(row_index, column_index).text = value
        for row in table.rows:
            for cell in row.cells:
                cell.text_frame.paragraphs[0].font.size = Pt(11)
    else:
        add_bullets(
            slide,
            [
                "No measured result file is available.",
                "Run python -m evaluation.run_all --videos videos.yaml after replacing placeholders and configuring .env.",
                "The evaluation runner regenerates this slide deck from results/results.csv.",
            ],
            top=2.0,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("docs/aat_presentation.pptx"))
    parser.add_argument("--results", type=Path, default=Path("results/results.csv"))
    arguments = parser.parse_args()
    print(build_presentation(arguments.output, arguments.results))


if __name__ == "__main__":
    main()
